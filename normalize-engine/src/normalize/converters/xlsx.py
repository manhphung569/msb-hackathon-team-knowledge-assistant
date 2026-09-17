from __future__ import annotations

from pathlib import Path

import openpyxl


def _sheet_to_markdown(ws) -> str:
    max_col = ws.max_column or 0
    rows = []
    for row in ws.iter_rows(values_only=True):
        if any(c is not None and str(c).strip() for c in row):
            rows.append(["" if c is None else str(c).strip() for c in row])
    if not rows:
        return "(sheet trống)"

    def pad(r: list[str]) -> list[str]:
        return r + [""] * (max_col - len(r))

    header = pad(rows[0])
    lines = ["| " + " | ".join(header) + " |", "| " + " | ".join(["---"] * max_col) + " |"]
    for row in rows[1:]:
        lines.append("| " + " | ".join(pad(row)) + " |")
    return "\n".join(lines)


def convert(path: Path) -> str:
    wb = openpyxl.load_workbook(str(path), data_only=True)
    parts = [f"## Sheet: {name}\n\n{_sheet_to_markdown(wb[name])}" for name in wb.sheetnames]
    return "\n\n".join(parts)
