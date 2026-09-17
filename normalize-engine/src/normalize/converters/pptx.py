from __future__ import annotations

from pathlib import Path

from pptx import Presentation


def _table_to_markdown(table) -> str:
    rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
    if not rows:
        return ""
    header = rows[0]
    lines = ["| " + " | ".join(header) + " |", "| " + " | ".join(["---"] * len(header)) + " |"]
    for row in rows[1:]:
        lines.append("| " + " | ".join(row) + " |")
    return "\n".join(lines)


def convert(path: Path) -> str:
    prs = Presentation(str(path))
    parts: list[str] = []
    for i, slide in enumerate(prs.slides, start=1):
        pieces: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    pieces.append(text)
            if shape.has_table:
                md = _table_to_markdown(shape.table)
                if md:
                    pieces.append(md)
        body = "\n\n".join(pieces) if pieces else "(slide không có text)"
        parts.append(f"## Slide {i}\n\n{body}")
    return "\n\n".join(parts)
