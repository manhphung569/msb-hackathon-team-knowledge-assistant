"""Extracted verbatim from pdlc-core/adlc/tools/dashboard/server.py (source line numbers noted per block). RAW EXTRACTION — not yet import-fixed or adapted; see workbench/EXTRACTION_NOTES.md."""

import os
import re
from pathlib import Path

# ---- source line 212 (_detect_whisper_hallucination) ----
def _detect_whisper_hallucination(text: str) -> str | None:
    normalized = re.sub(r"\s+", " ", (text or "").strip().lower())
    if not normalized:
        return None

    canned_phrases = (
        "cảm ơn bạn đã xem",
        "hẹn gặp lại",
        "thanks for watching",
        "please subscribe",
        "like and subscribe",
        "see you next time",
        "đăng ký kênh",
        "đăng kí kênh",
        "đăng ký cho kênh",
        "đăng kí cho kênh",
        "không bỏ lỡ những video",
    )
    if any(phrase in normalized for phrase in canned_phrases):
        return "hallucination:boilerplate"

    phrases = [
        part.strip(" -")
        for part in re.split(r"[,;.!?\n]+", normalized)
        if part.strip(" -")
    ]
    if len(phrases) >= 3:
        counts: dict[str, int] = {}
        for phrase in phrases:
            if len(phrase) < 8:
                continue
            counts[phrase] = counts.get(phrase, 0) + 1
        if counts:
            dominant_phrase, dominant_count = max(counts.items(), key=lambda item: item[1])
            repeated_chars = dominant_count * len(dominant_phrase)
            if dominant_count >= 3 and repeated_chars >= max(48, int(len(normalized) * 0.55)):
                return "hallucination:repetition"

    tokens = re.findall(r"\w+", normalized)
    if len(normalized) >= 80 and len(tokens) >= 12:
        unique_ratio = len(set(tokens)) / len(tokens)
        if unique_ratio <= 0.45:
            return "hallucination:low-entropy"

    return None


# ---- source line 2119 (_TEXT_EXTS_SET) ----
_TEXT_EXTS_SET = {
    ".txt",".md",".markdown",".json",".yaml",".yml",".xml",".csv",
    ".py",".java",".js",".ts",".jsx",".tsx",".html",".htm",".css",
    ".scss",".sql",".sh",".bash",".go",".rs",".cpp",".c",".h",
    ".php",".rb",".kt",".swift",".dart",".vue",".svelte",".toml",
    ".ini",".env",".log",".conf",".config",".properties",".gradle",
    ".pom",".lock",".rtf",
}


# ---- source line 2127 (_IMG_EXTS_SET) ----
_IMG_EXTS_SET  = {".png",".jpg",".jpeg",".gif",".webp",".svg",".bmp",".ico"}


# ---- source line 2128 (_ARCHIVE_EXTS) ----
_ARCHIVE_EXTS  = {".zip",".rar",".7z",".tar",".gz",".bz2",".xz",".tgz",".tar.gz",".tar.bz2",".tar.xz"}


# ---- source line 2131 (_ole_extract_text) ----
def _ole_extract_text(file_data: bytes) -> str:
    """Heuristic text extraction from OLE (Compound Document) binary files (.doc, .ppt).
    Not perfect — extracts readable Unicode sequences from raw streams.
    """
    import olefile, re as _re, io as _io
    texts: list[str] = []
    try:
        ole = olefile.OleFileIO(_io.BytesIO(file_data))
        # Streams that may contain text
        candidates = [n for n in ole.listdir() if isinstance(n, list)]
        priority = ["WordDocument", "PowerPoint Document", "1Table", "0Table", "SummaryInformation"]
        stream_names = sorted(
            ["/".join(p) for p in candidates],
            key=lambda s: (0 if any(s.startswith(p) for p in priority) else 1, s),
        )
        for sn in stream_names[:8]:
            try:
                raw = ole.openstream(sn).read()
                # Try UTF-16LE (modern Office internal encoding)
                decoded = raw.decode("utf-16-le", errors="ignore")
                chunks = _re.findall(r"[\x20-\x7EÀ-ɏḀ-ỿЀ-ӿ]{4,}", decoded)
                chunk_text = " ".join(chunks)
                if len(chunk_text.strip()) > 30:
                    texts.append(chunk_text)
            except Exception:
                pass
        ole.close()
    except Exception:
        pass
    combined = "\n".join(texts)
    # Deduplicate repeated whitespace
    import re as _re2
    return _re2.sub(r" {3,}", "  ", _re2.sub(r"\n{3,}", "\n\n", combined)).strip()


# ---- source line 2166 (_extract_legacy_office_via_com) ----
def _extract_legacy_office_via_com(file_data: bytes, ext: str) -> str:
    """Use local Office COM automation on Windows for real legacy .doc/.ppt files.

    This is more reliable than raw OLE heuristics for Word/PowerPoint 97-2003 binaries.
    Returns an empty string when COM is unavailable or extraction fails.
    """
    import base64 as _base64
    import os as _os
    import subprocess as _subprocess
    import tempfile as _tempfile
    import textwrap as _textwrap

    if _os.name != "nt" or ext not in (".doc", ".ppt"):
        return ""

    with _tempfile.NamedTemporaryFile(delete=False, suffix=ext) as tmp:
        tmp.write(file_data)
        tmp_path = tmp.name

    ps_path = tmp_path.replace("'", "''")
    if ext == ".doc":
        script = _textwrap.dedent(
            f"""
            $ErrorActionPreference = 'Stop'
            $docPath = '{ps_path}'
            $word = $null
            $doc = $null
            try {{
                $word = New-Object -ComObject Word.Application
                $word.Visible = $false
                $doc = $word.Documents.Open([string]$docPath, $false, $true)
                $text = [string]$doc.Content.Text
                $bytes = [System.Text.Encoding]::UTF8.GetBytes($text)
                [Console]::Out.Write([Convert]::ToBase64String($bytes))
            }} finally {{
                if ($doc) {{ $doc.Close($false) | Out-Null }}
                if ($word) {{ $word.Quit() }}
            }}
            """
        ).strip()
    else:
        script = _textwrap.dedent(
            f"""
            $ErrorActionPreference = 'Stop'
            $pptPath = '{ps_path}'
            $app = $null
            $presentation = $null
            $lines = New-Object System.Collections.Generic.List[string]
            try {{
                $app = New-Object -ComObject PowerPoint.Application
                $presentation = $app.Presentations.Open($pptPath, $false, $true, $false)
                foreach ($slide in $presentation.Slides) {{
                    $slideLines = New-Object System.Collections.Generic.List[string]
                    foreach ($shape in $slide.Shapes) {{
                        try {{
                            if ($shape.HasTextFrame -and $shape.TextFrame.HasText) {{
                                $shapeText = [string]$shape.TextFrame.TextRange.Text
                                if ($shapeText) {{
                                    foreach ($line in ($shapeText -split "`r?`n")) {{
                                        $trimmed = $line.Trim()
                                        if ($trimmed) {{ $slideLines.Add($trimmed) }}
                                    }}
                                }}
                            }}
                        }} catch {{}}
                    }}
                    if ($slideLines.Count -gt 0) {{
                        $lines.Add("=== Slide $($slide.SlideIndex) ===")
                        foreach ($line in $slideLines) {{ $lines.Add($line) }}
                    }}
                }}
                $text = $lines -join "`n"
                $bytes = [System.Text.Encoding]::UTF8.GetBytes($text)
                [Console]::Out.Write([Convert]::ToBase64String($bytes))
            }} finally {{
                if ($presentation) {{ $presentation.Close() }}
                if ($app) {{ $app.Quit() }}
            }}
            """
        ).strip()

    startupinfo = None
    if hasattr(_subprocess, "STARTUPINFO"):
        startupinfo = _subprocess.STARTUPINFO()
        startupinfo.dwFlags |= _subprocess.STARTF_USESHOWWINDOW
        startupinfo.wShowWindow = 0

    try:
        completed = _subprocess.run(
            [
                "powershell.exe",
                "-NoLogo",
                "-NoProfile",
                "-NonInteractive",
                "-ExecutionPolicy",
                "Bypass",
                "-Command",
                script,
            ],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=60,
            startupinfo=startupinfo,
            check=False,
        )
        payload = (completed.stdout or "").strip()
        if completed.returncode != 0 or not payload:
            return ""
        return _base64.b64decode(payload).decode("utf-8", errors="replace").strip()
    except Exception:
        return ""
    finally:
        try:
            _os.remove(tmp_path)
        except OSError:
            pass


# ---- source line 2286 (_extract_archive_text) ----
def _extract_archive_text(file_data: bytes, ext: str, filename: str) -> dict:
    """Extract text content from archive files (zip, 7z, rar, tar variants).
    Returns text listing + content of readable files inside.
    """
    import io as _io, re as _re

    TEXT_THRESHOLD = 512 * 1024   # 512 KB per file inside archive
    MAX_FILES_CONTENT = 10        # max files to read content from
    MAX_TOTAL_CHARS = 18000

    lines: list[str] = []
    file_list: list[str] = []
    content_parts: list[str] = []
    content_count = 0
    total_chars = 0

    def _is_readable(name: str) -> bool:
        e = Path(name).suffix.lower()
        return e in _TEXT_EXTS_SET or e in {".php", ".rb", ".go"}

    def _is_img(name: str) -> bool:
        return Path(name).suffix.lower() in _IMG_EXTS_SET

    def _add_content(name: str, data: bytes) -> None:
        nonlocal content_count, total_chars
        if content_count >= MAX_FILES_CONTENT or total_chars >= MAX_TOTAL_CHARS:
            return
        try:
            text = data[:TEXT_THRESHOLD].decode("utf-8", errors="replace")
            truncated = len(data) > TEXT_THRESHOLD
            note = " (đã cắt bớt)" if truncated else ""
            snippet = text[:3000]
            content_parts.append(f"--- {name}{note} ---\n{snippet}")
            content_count += 1
            total_chars += len(snippet)
        except Exception:
            pass

    # ── ZIP ──────────────────────────────────────────────────────────────────
    if ext == ".zip":
        import zipfile as _zf
        try:
            with _zf.ZipFile(_io.BytesIO(file_data)) as z:
                all_names = [i.filename for i in z.infolist() if not i.is_dir()]
                file_list = all_names
                for info in z.infolist():
                    if info.is_dir():
                        continue
                    name = info.filename
                    if _is_readable(name) and info.file_size < TEXT_THRESHOLD:
                        try:
                            _add_content(name, z.read(name))
                        except Exception:
                            pass
        except Exception as e:
            raise ValueError(f"Không đọc được ZIP: {e}")

    # ── 7Z ──────────────────────────────────────────────────────────────────
    elif ext == ".7z":
        import py7zr, tempfile, shutil, os as _os
        tmp_dir = tempfile.mkdtemp(prefix="vault7z_")
        try:
            with py7zr.SevenZipFile(_io.BytesIO(file_data), mode="r") as z:
                all_info = z.list()
                file_list = [f.filename for f in all_info if not f.is_directory]
                targets = [f.filename for f in all_info
                           if not f.is_directory and _is_readable(f.filename)
                           and (f.uncompressed or 0) < TEXT_THRESHOLD][:MAX_FILES_CONTENT]
                if targets:
                    z.extract(targets=targets, path=tmp_dir)
            # Read extracted files
            for name in targets:
                fpath = _os.path.join(tmp_dir, name.replace("/", _os.sep))
                if _os.path.isfile(fpath):
                    with open(fpath, "rb") as fh:
                        _add_content(name, fh.read())
        except Exception as e:
            raise ValueError(f"Không đọc được 7Z: {e}")
        finally:
            shutil.rmtree(tmp_dir, ignore_errors=True)

    # ── RAR ──────────────────────────────────────────────────────────────────
    elif ext == ".rar":
        import rarfile as _rf
        try:
            with _rf.RarFile(_io.BytesIO(file_data)) as z:
                all_info = z.infolist()
                file_list = [i.filename for i in all_info if not i.is_dir()]
                # Try to extract (needs unrar binary; fail gracefully)
                try:
                    for info in all_info:
                        if info.is_dir():
                            continue
                        if _is_readable(info.filename) and info.file_size < TEXT_THRESHOLD:
                            _add_content(info.filename, z.read(info.filename))
                except Exception:
                    content_parts.append("[Ghi chú: cần cài unrar để đọc nội dung RAR. "
                                         "Danh sách file bên trên đã được liệt kê.]")
        except Exception as e:
            raise ValueError(f"Không đọc được RAR: {e}")

    # ── TAR variants ─────────────────────────────────────────────────────────
    elif ext in (".tar", ".gz", ".bz2", ".xz", ".tgz"):
        import tarfile as _tf
        mode_map = {".tar": "r:", ".gz": "r:gz", ".tgz": "r:gz",
                    ".bz2": "r:bz2", ".xz": "r:xz"}
        mode = mode_map.get(ext, "r:*")
        try:
            with _tf.open(fileobj=_io.BytesIO(file_data), mode=mode) as z:
                members = [m for m in z.getmembers() if m.isfile()]
                file_list = [m.name for m in members]
                for m in members:
                    if _is_readable(m.name) and m.size < TEXT_THRESHOLD:
                        try:
                            f = z.extractfile(m)
                            if f:
                                _add_content(m.name, f.read())
                        except Exception:
                            pass
        except Exception as e:
            raise ValueError(f"Không đọc được archive: {e}")
    else:
        raise ValueError(f"Định dạng archive '{ext}' chưa hỗ trợ")

    # ── Build summary ────────────────────────────────────────────────────────
    imgs = [n for n in file_list if _is_img(n)]
    txts = [n for n in file_list if _is_readable(n)]
    bins = [n for n in file_list if not _is_readable(n) and not _is_img(n)]

    summary = [f"=== Archive: {filename} ({len(file_list)} files) ==="]
    if txts:
        summary.append(f"📄 Text/Code ({len(txts)}): " + ", ".join(txts[:20]))
    if imgs:
        summary.append(f"🖼 Ảnh ({len(imgs)}): " + ", ".join(imgs[:10]))
    if bins:
        summary.append(f"📦 Binary ({len(bins)}): " + ", ".join(bins[:10]))

    all_text = "\n".join(summary)
    if content_parts:
        all_text += "\n\n=== Nội dung file ===\n\n" + "\n\n".join(content_parts)

    meta = {"files": len(file_list), "text_files": len(txts),
            "images": len(imgs), "binary": len(bins)}
    return {"text": all_text, "meta": meta}


# ---- source line 2432 (_decode_text_bytes) ----
def _decode_text_bytes(file_data: bytes) -> str:
    """Best-effort decode for text-like uploads using common encodings."""
    for encoding in ("utf-8-sig", "utf-16", "utf-16-le", "utf-16-be", "cp1252", "latin-1"):
        try:
            text = file_data.decode(encoding)
        except Exception:
            continue
        if "\x00" in text[:512]:
            continue
        return text
    return file_data.decode("utf-8", errors="replace")


# ---- source line 2445 (_extract_markup_text) ----
def _extract_markup_text(file_data: bytes) -> dict:
    """Extract readable text from markup-heavy formats like HTML/XML."""
    import html as _html
    import re as _re

    raw = _decode_text_bytes(file_data)
    text = _re.sub(r"(?is)<(script|style|noscript).*?>.*?</\\1>", " ", raw)
    text = _re.sub(r"(?is)<!--.*?-->", " ", text)
    text = _re.sub(r"(?s)<[^>]+>", " ", text)
    text = _html.unescape(text)
    text = _re.sub(r"\s+", " ", text).strip()
    return {"text": text, "meta": {"note": "markup-text"}}


# ---- source line 2460 (_extract_document_text) ----
def _extract_document_text(file_data: bytes, ext: str, filename: str = "") -> dict:
    """Extract plain text from binary document formats.
    Returns dict with 'text' key and optional 'meta' dict.
    Raises ValueError for unsupported formats.
    """
    import io as _io

    # ── Plain text / markup formats ─────────────────────────────────────────
    if ext in (
        ".txt", ".md", ".markdown", ".json", ".yaml", ".yml", ".xml", ".csv",
        ".html", ".htm", ".css", ".scss", ".sql", ".sh", ".bash", ".go", ".rs",
        ".cpp", ".c", ".h", ".py", ".java", ".js", ".ts", ".jsx", ".tsx", ".php",
        ".rb", ".kt", ".swift", ".dart", ".vue", ".svelte", ".toml", ".ini", ".env",
        ".log", ".conf", ".config", ".properties", ".gradle", ".pom", ".lock",
    ):
        if ext in (".html", ".htm", ".xml"):
            return _extract_markup_text(file_data)
        return {"text": _decode_text_bytes(file_data)}

    # ── Archive formats ──────────────────────────────────────────────────────
    if ext in (".zip", ".7z", ".rar", ".tar", ".gz", ".bz2", ".xz", ".tgz"):
        return _extract_archive_text(file_data, ext, filename)

    # ── PDF ──────────────────────────────────────────────────────────────────
    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(_io.BytesIO(file_data))
        pages = []
        for i, page in enumerate(reader.pages, 1):
            t = (page.extract_text() or "").strip()
            if t:
                pages.append(f"=== Trang {i} ===\n{t}")
        return {"text": "\n\n".join(pages), "meta": {"pages": len(reader.pages)}}

    # ── DOCX ─────────────────────────────────────────────────────────────────
    elif ext in (".docx", ".docm", ".dotx", ".dotm"):
        import docx as _docx
        doc = _docx.Document(_io.BytesIO(file_data))
        parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(c.text.strip() for c in row.cells)
                if row_text.strip(" |"):
                    parts.append(row_text)
        return {"text": "\n".join(parts), "meta": {"paragraphs": len(doc.paragraphs)}}

    # ── DOC (binary) — olefile heuristic ─────────────────────────────────────
    elif ext == ".doc":
        # Try as docx first (some .doc files are actually docx-renamed)
        try:
            import docx as _docx
            doc = _docx.Document(_io.BytesIO(file_data))
            parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            if parts:
                return {"text": "\n".join(parts), "meta": {"note": "Đọc như DOCX"}}
        except Exception:
            pass
        com_text = _extract_legacy_office_via_com(file_data, ext)
        if len(com_text.strip()) >= 20:
            return {"text": com_text, "meta": {"note": "Word COM"}}
        # Fallback: OLE heuristic
        text = _ole_extract_text(file_data)
        if len(text.strip()) < 20:
            raise ValueError(
                "Không thể đọc .doc này. "
                "Hãy mở bằng Word/LibreOffice và lưu lại dưới dạng .docx"
            )
        return {"text": text, "meta": {"note": "Heuristic OLE — có thể thiếu định dạng"}}

    # ── XLSX / XLSM ──────────────────────────────────────────────────────────
    elif ext in (".xlsx", ".xlsm", ".xltx", ".xltm"):
        import openpyxl as _xl
        wb = _xl.load_workbook(_io.BytesIO(file_data), data_only=True, read_only=True)
        lines = []
        for sheet in wb.worksheets:
            lines.append(f"=== Sheet: {sheet.title} ===")
            for row in sheet.iter_rows(values_only=True):
                cells = [str(v) if v is not None else "" for v in row]
                if any(c.strip() for c in cells):
                    lines.append("\t".join(cells))
        meta = {"sheets": len(wb.sheetnames)}
        wb.close()
        return {"text": "\n".join(lines), "meta": meta}

    # ── XLS (binary Excel 97-2003) ────────────────────────────────────────────
    elif ext == ".xls":
        import xlrd
        wb = xlrd.open_workbook(file_contents=file_data)
        lines = []
        for i in range(wb.nsheets):
            ws = wb.sheet_by_index(i)
            lines.append(f"=== Sheet: {ws.name} ===")
            for r in range(ws.nrows):
                row_vals = []
                for c in range(ws.ncols):
                    cell = ws.cell(r, c)
                    row_vals.append(str(cell.value) if cell.value != "" else "")
                row_str = "\t".join(row_vals)
                if row_str.strip():
                    lines.append(row_str)
        return {"text": "\n".join(lines), "meta": {"sheets": wb.nsheets, "note": "Excel 97-2003"}}

    # ── PPTX / PPTM ──────────────────────────────────────────────────────────
    elif ext in (".pptx", ".pptm", ".ppsx", ".ppsm", ".potx", ".potm"):
        from pptx import Presentation as _Prs
        prs = _Prs(_io.BytesIO(file_data))
        lines = []
        for i, slide in enumerate(prs.slides, 1):
            slide_texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        t = para.text.strip()
                        if t:
                            slide_texts.append(t)
            if slide_texts:
                lines.append(f"=== Slide {i} ===")
                lines.extend(slide_texts)
        return {"text": "\n".join(lines), "meta": {"slides": len(prs.slides)}}

    # ── PPT (binary) — olefile heuristic ─────────────────────────────────────
    elif ext == ".ppt":
        com_text = _extract_legacy_office_via_com(file_data, ext)
        if len(com_text.strip()) >= 20:
            return {"text": com_text, "meta": {"note": "PowerPoint COM"}}
        text = _ole_extract_text(file_data)
        if len(text.strip()) < 20:
            raise ValueError(
                "Không thể đọc .ppt này. "
                "Hãy mở bằng PowerPoint/LibreOffice và lưu lại dưới dạng .pptx"
            )
        return {"text": text, "meta": {"note": "Heuristic OLE — có thể thiếu định dạng"}}

    # ── RTF ──────────────────────────────────────────────────────────────────
    elif ext == ".rtf":
        import re as _re
        text = file_data.decode("latin-1", errors="replace")
        text = _re.sub(r"\\\*\\[a-z]+[^}]*", "", text)
        text = _re.sub(r"\\u-?\d+ ?", lambda m: " ", text)  # Unicode escapes
        text = _re.sub(r"\\[a-z]+\d* ?", "", text)
        text = _re.sub(r"[{}]", "", text)
        text = _re.sub(r"\s+", " ", text).strip()
        return {"text": text}

    # ── LibreOffice ODF formats ───────────────────────────────────────────────
    elif ext in (".odt", ".ods", ".odp", ".odg"):
        import zipfile as _zf, re as _re
        with _zf.ZipFile(_io.BytesIO(file_data)) as z:
            content = z.read("content.xml").decode("utf-8", errors="replace")
        text = _re.sub(r"<[^>]+>", " ", content)
        text = _re.sub(r"\s+", " ", text).strip()
        return {"text": text}

    else:
        raise ValueError(f"Định dạng '{ext}' chưa được hỗ trợ")


class ExtractionMixin:
    # ---- source line 5359 (Handler._handle_vault_extract_file) ----
    def _handle_vault_extract_file(self):
        """POST /vault/extract-file
        Headers: X-Filename: <url-encoded filename>
        Body:    raw binary file content
        Returns: {text, truncated, filename, pages/sheets/slides?}
        """
        import io as _io
        from urllib.parse import unquote as _unquote

        filename = _unquote(self.headers.get("X-Filename", "file"))
        clen = int(self.headers.get("Content-Length", 0))
        if clen > 50 * 1024 * 1024:  # 50 MB cap
            return self._send_json({"error": "File quá lớn (max 50 MB)"}, 400)
        file_data = self.rfile.read(clen)
        ext = Path(filename).suffix.lower()

        try:
            result = _extract_document_text(file_data, ext, filename)
        except ModuleNotFoundError as e:
            return self._send_json({
                "error": f"Thiếu package đọc file: {e.name}. Hãy cài theo workbench/requirements.txt"
            }, 400)
        except Exception as e:
            return self._send_json({"error": str(e)}, 400)

        MAX = 20000
        text = result["text"]
        truncated = len(text) > MAX
        out = {"filename": filename, "ext": ext, "text": text[:MAX], "truncated": truncated}
        if "meta" in result:
            out["meta"] = result["meta"]
        return self._send_json(out)

    # ---- source line 8358 (Handler._handle_whisper_key) ----
    def _handle_whisper_key(self):
        """GET /vault/whisper-key — return OPENAI_API_KEY for client-side Whisper calls (localhost only)."""
        key = os.environ.get("OPENAI_API_KEY", "")
        self._send_json({"key": key})

    # ---- source line 8363 (Handler._handle_whisper_transcribe) ----
    def _handle_whisper_transcribe(self):
        """POST /vault/whisper-transcribe — proxy to OpenAI Whisper using requests library."""
        import traceback
        try:
            self._do_whisper_transcribe()
        except Exception as e:
            traceback.print_exc()
            try:
                self._send_json({"error": str(e)}, 500)
            except Exception:
                pass

    # ---- source line 8375 (Handler._do_whisper_transcribe) ----
    def _do_whisper_transcribe(self):
        import requests as _req

        key = os.environ.get("OPENAI_API_KEY", "")
        if not key:
            self._send_json({"error": "No OPENAI_API_KEY configured"}, 400); return

        length      = int(self.headers.get("Content-Length", 0))
        audio_bytes = self.rfile.read(length)
        mime        = self.headers.get("Content-Type", "audio/webm").split(";")[0].strip()
        ext         = "ogg" if "ogg" in mime else "webm"
        print(f"[Whisper] received {len(audio_bytes)} bytes  mime={mime}", flush=True)

        resp = _req.post(
            "https://api.openai.com/v1/audio/transcriptions",
            headers={"Authorization": f"Bearer {key}"},
            files={"file": (f"audio.{ext}", audio_bytes, mime)},
            data={"model": "whisper-1", "language": "vi",
                  "response_format": "verbose_json", "temperature": "0"},
            timeout=45,
        )
        print(f"[Whisper] OpenAI status={resp.status_code}", flush=True)
        if not resp.ok:
            print(f"[Whisper] error body: {resp.text[:300]}", flush=True)
            self._send_json({"error": f"Whisper {resp.status_code}: {resp.text[:200]}"}, 502); return

        result   = resp.json()
        text     = result.get("text", "").strip()
        segments = result.get("segments", [])
        avg_nsp  = (sum(s.get("no_speech_prob", 0) for s in segments) / len(segments)) if segments else 0
        print(f"[Whisper] text={repr(text[:100])}  nsp={avg_nsp:.3f}", flush=True)

        if avg_nsp > 0.80:
            self._send_json({"text": "", "filtered": True, "reason": f"nsp={avg_nsp:.2f}"}); return

        hallucination_reason = _detect_whisper_hallucination(text)
        if hallucination_reason:
            print(f"[Whisper] hallucination filtered: {repr(text)}", flush=True)
            self._send_json({"text": "", "filtered": True, "reason": hallucination_reason}); return

        self._send_json({"text": text})

